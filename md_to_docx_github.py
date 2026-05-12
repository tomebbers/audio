"""
Converts Markdown to a Word document with GitHub-preview styling using python-docx.
Handles: headings, paragraphs, tables, bullet/numbered lists, bold/italic/code,
         horizontal rules, blockquotes.
"""
import os, re
os.chdir('/home/user/audio')

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import xml.etree.ElementTree as ET

# ── colours matching GitHub preview ──────────────────────────────────────────
C_BLACK   = RGBColor(0x24, 0x29, 0x2e)
C_GRAY    = RGBColor(0x6a, 0x73, 0x7d)
C_BLUE    = RGBColor(0x03, 0x66, 0xd6)
C_H_BLUE  = RGBColor(0x1f, 0x49, 0x7d)   # dark blue for H1/H2
C_RED     = RGBColor(0xd7, 0x3a, 0x49)
C_TBL_HDR = RGBColor(0xf6, 0xf8, 0xfa)

FONT_BODY = "Calibri"
FONT_MONO = "Courier New"

# ── helpers ───────────────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color="F6F8FA"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_bottom_border(paragraph):
    """Add a bottom border line under a paragraph (GitHub H1/H2 style)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'EAECEF')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_left_border(paragraph, color="DFE2E5"):
    """Add a left border (blockquote style)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '12')
    left.set(qn('w:space'), '4')
    left.set(qn('w:color'), color)
    pBdr.append(left)
    pPr.append(pBdr)

def set_para_spacing(para, before=0, after=8):
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:before'), str(int(before * 20)))
    spacing.set(qn('w:after'), str(int(after * 20)))
    pPr.append(spacing)

def apply_inline(run, bold=False, italic=False, code=False,
                 color=None, font_size=None):
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if code:
        run.font.name = FONT_MONO
        run.font.size = Pt(10)
        # light gray shading via rPr highlight isn't available directly;
        # set color to dark gray instead
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2e)
    if color:
        run.font.color.rgb = color
    if font_size:
        run.font.size = Pt(font_size)

# ── inline Markdown parser ────────────────────────────────────────────────────
# returns list of (text, bold, italic, code) tuples

INLINE_RE = re.compile(
    r'(?P<bold_it>\*{3}(?P<bi_text>.+?)\*{3})'
    r'|(?P<bold>\*{2}(?P<b_text>.+?)\*{2})'
    r'|(?P<italic>\*(?P<i_text>.+?)\*)'
    r'|(?P<code>`(?P<c_text>[^`]+?)`)'
    r'|(?P<text>[^`*]+)'
    r'|(?P<star>[*`]+)',
    re.DOTALL
)

def parse_inline(text):
    """Return list of (text, bold, italic, code)."""
    parts = []
    for m in INLINE_RE.finditer(text):
        if m.group('bold_it'):
            parts.append((m.group('bi_text'), True, True, False))
        elif m.group('bold'):
            parts.append((m.group('b_text'), True, False, False))
        elif m.group('italic'):
            parts.append((m.group('i_text'), False, True, False))
        elif m.group('code'):
            parts.append((m.group('c_text'), False, False, True))
        elif m.group('text'):
            parts.append((m.group('text'), False, False, False))
        elif m.group('star'):
            parts.append((m.group('star'), False, False, False))
    return parts

def add_inline_text(para, text, base_bold=False, base_italic=False,
                    base_color=None, base_size=None):
    for segment, bold, italic, code in parse_inline(text):
        run = para.add_run(segment)
        run.font.name = FONT_BODY
        if base_size:
            run.font.size = Pt(base_size)
        apply_inline(run, bold=bold or base_bold,
                     italic=italic or base_italic,
                     code=code, color=base_color)

# ── document builder ──────────────────────────────────────────────────────────

class GithubDocx:
    def __init__(self):
        self.doc = Document()
        self._setup_document()

    def _setup_document(self):
        # Page margins
        for section in self.doc.sections:
            section.left_margin   = Cm(2.5)
            section.right_margin  = Cm(2.5)
            section.top_margin    = Cm(2.5)
            section.bottom_margin = Cm(2.5)

        # Default paragraph font
        self.doc.styles['Normal'].font.name = FONT_BODY
        self.doc.styles['Normal'].font.size = Pt(10.5)
        self.doc.styles['Normal'].font.color.rgb = C_BLACK

    def heading(self, text, level):
        sizes   = {1: 20, 2: 15, 3: 13, 4: 11, 5: 10, 6: 10}
        colors  = {1: C_H_BLUE, 2: C_H_BLUE, 3: C_H_BLUE,
                   4: C_BLACK, 5: C_BLACK, 6: C_GRAY}
        before  = {1: 24, 2: 20, 3: 16, 4: 12, 5: 8, 6: 8}

        para = self.doc.add_paragraph()
        set_para_spacing(para, before=before.get(level, 12), after=6)

        run = para.add_run(text)
        run.bold = True
        run.font.name = FONT_BODY
        run.font.size = Pt(sizes.get(level, 11))
        run.font.color.rgb = colors.get(level, C_BLACK)

        if level == 1:
            add_bottom_border(para)
            run.font.all_caps = True
        if level == 2:
            add_bottom_border(para)

        return para

    def paragraph(self, text, indent=0, color=None, size=None):
        if not text.strip():
            return
        para = self.doc.add_paragraph()
        set_para_spacing(para, before=0, after=6)
        if indent:
            para.paragraph_format.left_indent = Cm(indent)
        add_inline_text(para, text, base_color=color, base_size=size)
        return para

    def bullet(self, text, level=0):
        para = self.doc.add_paragraph(style='List Bullet')
        para.paragraph_format.left_indent  = Cm(0.5 + level * 0.5)
        para.paragraph_format.first_line_indent = Cm(-0.3)
        set_para_spacing(para, before=0, after=2)
        add_inline_text(para, text)
        return para

    def numbered(self, text, level=0):
        para = self.doc.add_paragraph(style='List Number')
        para.paragraph_format.left_indent  = Cm(0.5 + level * 0.5)
        para.paragraph_format.first_line_indent = Cm(-0.3)
        set_para_spacing(para, before=0, after=2)
        add_inline_text(para, text)
        return para

    def blockquote(self, text):
        para = self.doc.add_paragraph()
        set_para_spacing(para, before=0, after=6)
        para.paragraph_format.left_indent = Cm(1.0)
        add_left_border(para, 'DFE2E5')
        add_inline_text(para, text, base_color=C_GRAY)
        return para

    def horizontal_rule(self):
        para = self.doc.add_paragraph()
        set_para_spacing(para, before=12, after=12)
        pPr = para._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '6')
        bottom.set(qn('w:space'), '1')
        bottom.set(qn('w:color'), 'E1E4E8')
        pBdr.append(bottom)
        pPr.append(pBdr)
        return para

    def table(self, headers, rows):
        col_count = max(len(headers), max((len(r) for r in rows), default=0))
        tbl = self.doc.add_table(rows=1 + len(rows), cols=col_count)
        tbl.style = 'Table Grid'

        # Header row
        hdr_row = tbl.rows[0]
        for i, h in enumerate(headers):
            cell = hdr_row.cells[i]
            set_cell_bg(cell, 'F6F8FA')
            para = cell.paragraphs[0]
            set_para_spacing(para, before=2, after=2)
            add_inline_text(para, h, base_bold=True, base_size=9.5)

        # Data rows
        for ri, row_data in enumerate(rows):
            word_row = tbl.rows[ri + 1]
            bg = 'F6F8FA' if ri % 2 == 1 else 'FFFFFF'
            for ci in range(col_count):
                cell = word_row.cells[ci]
                set_cell_bg(cell, bg)
                para = cell.paragraphs[0]
                set_para_spacing(para, before=2, after=2)
                text = row_data[ci] if ci < len(row_data) else ''
                add_inline_text(para, text, base_size=9.5)

        # Spacing after table
        after = self.doc.add_paragraph()
        set_para_spacing(after, before=0, after=6)
        return tbl

    def save(self, path):
        self.doc.save(path)
        print(f"Saved: {path}")


# ── Markdown line-by-line parser ──────────────────────────────────────────────

def parse_and_build(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    gdoc = GithubDocx()

    # State
    i = 0
    total = len(lines)

    def peek(n=0):
        idx = i + n
        return lines[idx].rstrip('\n') if idx < total else ''

    # list state
    list_stack = []   # stack of ('bullet'|'number', level)

    def flush_list():
        list_stack.clear()

    # table state
    in_table = False
    table_headers = []
    table_rows = []

    def flush_table():
        nonlocal in_table, table_headers, table_rows
        if table_headers or table_rows:
            gdoc.table(table_headers, table_rows)
        in_table = False
        table_headers = []
        table_rows = []

    while i < total:
        raw = lines[i].rstrip('\n')
        stripped = raw.strip()

        # ── blank line ──
        if stripped == '':
            flush_list()
            if in_table:
                flush_table()
            i += 1
            continue

        # ── horizontal rule ──
        if re.match(r'^---+$', stripped) or re.match(r'^\*\*\*+$', stripped):
            flush_list()
            if in_table:
                flush_table()
            gdoc.horizontal_rule()
            i += 1
            continue

        # ── ATX heading (# ## ### ...) ──
        m = re.match(r'^(#{1,6})\s+(.*)', stripped)
        if m:
            flush_list()
            if in_table:
                flush_table()
            level = len(m.group(1))
            text = m.group(2).strip()
            gdoc.heading(text, level)
            i += 1
            continue

        # ── table row ──
        if stripped.startswith('|') and stripped.endswith('|'):
            flush_list()
            cells = [c.strip() for c in stripped.strip('|').split('|')]

            # separator row (---|---|---)?
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                i += 1
                continue

            if not in_table:
                in_table = True
                table_headers = cells
            else:
                table_rows.append(cells)
            i += 1
            continue
        elif in_table:
            flush_table()

        # ── blockquote ──
        if stripped.startswith('>'):
            flush_list()
            text = re.sub(r'^>\s?', '', stripped)
            gdoc.blockquote(text)
            i += 1
            continue

        # ── unordered list ──
        m = re.match(r'^(\s*)[-*+]\s+(.*)', raw)
        if m:
            indent_len = len(m.group(1))
            level = indent_len // 2
            text = m.group(2)
            gdoc.bullet(text, level=level)
            if not list_stack or list_stack[-1] != ('bullet', level):
                list_stack.append(('bullet', level))
            i += 1
            continue

        # ── ordered list ──
        m = re.match(r'^(\s*)\d+[.)]\s+(.*)', raw)
        if m:
            indent_len = len(m.group(1))
            level = indent_len // 2
            text = m.group(2)
            gdoc.numbered(text, level=level)
            if not list_stack or list_stack[-1] != ('number', level):
                list_stack.append(('number', level))
            i += 1
            continue

        # ── fenced code block ──
        if stripped.startswith('```'):
            flush_list()
            if in_table:
                flush_table()
            i += 1
            code_lines = []
            while i < total and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            if code_lines:
                para = gdoc.doc.add_paragraph()
                set_para_spacing(para, before=4, after=4)
                para.paragraph_format.left_indent = Cm(0.5)
                run = para.add_run('\n'.join(code_lines))
                run.font.name = FONT_MONO
                run.font.size = Pt(9)
            i += 1
            continue

        # ── normal paragraph ──
        flush_list()
        # collect continuation lines
        text_parts = [stripped]
        i += 1
        while i < total:
            nxt = lines[i].rstrip('\n')
            nxt_s = nxt.strip()
            if (not nxt_s or nxt_s.startswith('#') or nxt_s.startswith('|')
                    or nxt_s.startswith('>') or nxt_s.startswith('```')
                    or re.match(r'^[-*+]\s', nxt_s)
                    or re.match(r'^\d+[.)]\s', nxt_s)
                    or re.match(r'^---+$', nxt_s)):
                break
            text_parts.append(nxt_s)
            i += 1
        gdoc.paragraph(' '.join(text_parts))

    if in_table:
        flush_table()

    gdoc.save(docx_path)


parse_and_build(
    '/home/user/audio/KNO_Richtlijnen_Uitgebreid.md',
    '/home/user/audio/KNO_Richtlijnen_Uitgebreid.docx'
)
